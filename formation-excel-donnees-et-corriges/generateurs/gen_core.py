# -*- coding: utf-8 -*-
"""Noyau de generation : tirage des transactions et utilitaires Excel.

Toutes les fonctions sont deterministes a graine fixee : relancer les
scripts reproduit exactement les memes fichiers.
"""

import datetime as dt
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import referentiel as ref

SORTIE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "donnees")

# Part de chaque categorie dans le nombre de transactions
POIDS_CATEGORIE = {
    "Électronique": 0.24,
    "Mobilier": 0.10,
    "Papeterie": 0.30,
    "Électroménager": 0.16,
    "Téléphonie": 0.20,
}


def dossier_sortie():
    if not os.path.isdir(SORTIE):
        os.makedirs(SORTIE)
    return SORTIE


def liste_mois(debut, fin):
    """Liste des couples (annee, mois) de debut a fin inclus."""
    mois = []
    a, m = debut.year, debut.month
    while (a, m) <= (fin.year, fin.month):
        mois.append((a, m))
        m += 1
        if m == 13:
            a, m = a + 1, 1
    return mois


def _jours_dans_mois(annee, mois):
    if mois == 12:
        suivant = dt.date(annee + 1, 1, 1)
    else:
        suivant = dt.date(annee, mois + 1, 1)
    return (suivant - dt.date(annee, mois, 1)).days


def _tirage_pondere(rng, elements, poids):
    total = sum(poids)
    seuil = rng.random() * total
    cumul = 0.0
    for element, p in zip(elements, poids):
        cumul += p
        if seuil <= cumul:
            return element
    return elements[-1]


def quantite_pour_prix(rng, prix, coef=1.0):
    """Quantite decroissante avec le prix : cree une correlation negative
    prix / quantite exploitable en nuage de points."""
    moyenne = 90.0 / (prix ** 0.62) * coef
    valeur = rng.gauss(moyenne, moyenne * 0.35)
    return max(1, min(80, int(round(valeur))))


def tirer_transactions(rng, debut, fin, nombre, croissance_annuelle=0.12,
                       categories=None, regions=None, vendeurs=None,
                       saison_region=None):
    """Tire `nombre` transactions entre `debut` et `fin`.

    `saison_region` est un profil mensuel PROPRE A LA REGION, independant
    des categories : dict {region: [12 coefficients]}. Il est necessaire
    des qu'un exercice interroge la saisonnalite d'une region, car
    amplifier la saisonnalite des categories ne produit pas de
    saisonnalite regionale : les categories culminent a des mois
    differents, et leurs pics s'annulent une fois agreges par region.

    Retourne une liste de dictionnaires triee par date croissante.
    """
    categories = categories or ref.CATEGORIES
    regions = regions or ref.REGIONS
    vendeurs = vendeurs or ref.VENDEURS

    mois = liste_mois(debut, fin)
    prod_par_cat = dict((c, ref.produits_de_categorie(c)) for c in categories)
    mag_par_region = {}
    vend_par_region = {}
    for r in regions:
        mag_par_region[r] = [m for m in ref.MAGASINS if m[3] == r]
        vend_par_region[r] = [v for v in vendeurs if v[1] == r]

    transactions = []
    for i in range(nombre):
        region = _tirage_pondere(
            rng, regions, [ref.POIDS_REGION[r] for r in regions])
        categorie = _tirage_pondere(
            rng, categories, [POIDS_CATEGORIE[c] for c in categories])

        amplitude = ref.AMPLITUDE_REGION[region]
        profil = (saison_region or {}).get(region)
        poids_mois = []
        for idx, (a, m) in enumerate(mois):
            saison = 1.0 + (ref.SAISON[categorie][m - 1] - 1.0) * amplitude
            if profil:
                saison *= profil[m - 1]
            tendance = (1.0 + croissance_annuelle) ** (idx / 12.0)
            poids_mois.append(max(0.05, saison) * tendance)
        annee, num_mois = _tirage_pondere(rng, mois, poids_mois)

        jour_max = _jours_dans_mois(annee, num_mois)
        borne_basse = debut.day if (annee, num_mois) == (debut.year, debut.month) else 1
        borne_haute = fin.day if (annee, num_mois) == (fin.year, fin.month) else jour_max
        if borne_haute < borne_basse:
            borne_haute = borne_basse
        jour = rng.randint(borne_basse, borne_haute)
        date = dt.date(annee, num_mois, jour)

        produit = rng.choice(prod_par_cat[categorie])
        magasin = rng.choice(mag_par_region[region])
        candidats = vend_par_region[region]
        vendeur = _tirage_pondere(rng, candidats, [v[3] for v in candidats])
        quantite = quantite_pour_prix(rng, produit[3], coef=vendeur[2])

        transactions.append({
            "date": date,
            "id_produit": produit[0],
            "produit": produit[1],
            "categorie": produit[2],
            "prix_unitaire": produit[3],
            "cout_unitaire": produit[4],
            "quantite": quantite,
            "id_magasin": magasin[0],
            "magasin": ref.MAGASINS_ACC[magasin[0]][0],
            "ville": ref.MAGASINS_ACC[magasin[0]][1],
            "region": ref.REGIONS_ACC[region],
            "region_brute": region,
            "vendeur": vendeur[0],
            "ca": round(produit[3] * quantite, 2),
            "cout_total": round(produit[4] * quantite, 2),
        })

    transactions.sort(key=lambda t: (t["date"], t["id_produit"]))
    for i, t in enumerate(transactions, start=1):
        t["id_vente"] = "V%05d" % i
    return transactions


# ----------------------------------------------------------------- Excel

GRIS = PatternFill("solid", fgColor="DDDDDD")
BLEU = PatternFill("solid", fgColor="1F3864")


def ecrire_feuille(ws, entetes, lignes, formats=None, nom_tableau=None,
                   largeurs=None, style_tableau="TableStyleMedium2"):
    """Ecrit un tableau (entetes + lignes) et le convertit en tableau
    structure Excel si `nom_tableau` est fourni."""
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = BLEU
        cellule.alignment = Alignment(horizontal="center", vertical="center",
                                      wrap_text=True)
    for ligne in lignes:
        ws.append(list(ligne))

    if formats:
        for col_idx, fmt in formats.items():
            if fmt is None:
                continue
            lettre = get_column_letter(col_idx)
            for ligne_idx in range(2, ws.max_row + 1):
                ws["%s%d" % (lettre, ligne_idx)].number_format = fmt

    if largeurs:
        for col_idx, largeur in largeurs.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = largeur
    else:
        for col_idx, entete in enumerate(entetes, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = max(
                12, min(30, len(str(entete)) + 4))

    ws.freeze_panes = "A2"

    if nom_tableau and ws.max_row > 1:
        ref_plage = "A1:%s%d" % (get_column_letter(len(entetes)), ws.max_row)
        tableau = Table(displayName=nom_tableau, ref=ref_plage)
        tableau.tableStyleInfo = TableStyleInfo(
            name=style_tableau, showRowStripes=True, showColumnStripes=False)
        ws.add_table(tableau)
    return ws


def nouveau_classeur(nom_premiere_feuille):
    wb = Workbook()
    wb.active.title = nom_premiere_feuille
    return wb


def enregistrer(wb, nom_fichier):
    chemin = os.path.join(dossier_sortie(), nom_fichier)
    wb.save(chemin)
    print("  ecrit : %s" % nom_fichier)
    return chemin


def rng_pour(cle):
    """Generateur aleatoire propre a un fichier : modifier un fichier
    ne decale pas les tirages des autres."""
    return random.Random("%s-%s" % (ref.SEED, cle))
