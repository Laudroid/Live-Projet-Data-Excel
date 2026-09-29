# -*- coding: utf-8 -*-
"""Suivi_Marges.xlsx -- fichier de travail du TP sur la robustesse.

Sert a l'enonce : Serie B / exercice 3 -- Optimiser la robustesse des
calculs (SIERREUR, gestion de #N/A et #DIV/0!).

Le fichier fournit les DONNEES et les EN-TETES des colonnes a calculer,
jamais les formules : les colonnes marquees "(a completer)" sont vides.

Pieges volontaires, chacun declenchant une erreur precise :
  * Saisies : 9 lignes portant un code hors catalogue        -> #N/A
  * Saisies : 7 quantites nulles ou vides (denominateur)     -> #DIV/0!
  * Saisies : 6 objectifs de CA nuls ou vides (denominateur) -> #DIV/0!
  * Catalogue : 2 produits dont le prix de vente vaut 0       -> #DIV/0!
  * Recherche : un code de test inexistant                    -> #N/A
"""

import datetime as dt

from openpyxl.styles import Alignment, Font, PatternFill

import gen_core as core
import referentiel as ref

# Trois references vendues mais jamais creees au catalogue,
# injectees sur 3 lignes chacune : exactement 9 lignes en #N/A.
CODES_INCONNUS = ["P777", "P812", "P950"]

JAUNE = PatternFill("solid", fgColor="FFF2CC")


def main():
    rng = core.rng_pour("suivi_marges")

    # ------------------------------------------------------- catalogue
    # Le catalogue est complet : les #N/A viennent uniquement des trois
    # codes inconnus injectes dans les saisies, ce qui rend le comptage
    # des erreurs previsible pour le formateur.
    catalogue = [[p[0], p[1], p[4], p[3]] for p in ref.PRODUITS]

    # deux produits en cours de tarification : prix de vente a zero
    for code in ["P105", "P304"]:
        for ligne in catalogue:
            if ligne[0] == code:
                ligne[3] = 0

    # ---------------------------------------------------------- saisies
    transactions = core.tirer_transactions(
        rng, dt.date(2024, 9, 1), dt.date(2024, 12, 31), 120,
        croissance_annuelle=0.0)

    saisies = []
    for i, t in enumerate(transactions, start=1):
        objectif = round(t["ca"] * rng.uniform(0.72, 1.28), 2)
        saisies.append([
            "L%03d" % i,
            t["date"],
            t["id_produit"],
            t["quantite"],
            t["ca"],
            objectif,
            None, None, None,
        ])

    indices = list(range(len(saisies)))
    rng.shuffle(indices)

    # codes absents du catalogue : 3 codes x 3 occurrences
    for k, i in enumerate(indices[0:9]):
        saisies[i][2] = CODES_INCONNUS[k % 3]

    # quantites nulles ou vides
    for k, i in enumerate(indices[9:16]):
        saisies[i][3] = 0 if k % 2 == 0 else None

    # objectifs nuls ou vides
    for k, i in enumerate(indices[16:22]):
        saisies[i][5] = 0 if k % 2 == 0 else None

    # ------------------------------------------------------- ecriture
    wb = core.nouveau_classeur("Catalogue")
    core.ecrire_feuille(
        wb["Catalogue"],
        ["Code_Produit", "Designation", "Prix_Achat", "Prix_Vente"],
        catalogue,
        formats={3: "# ##0.00 €", 4: "# ##0.00 €"},
        largeurs={1: 14, 2: 34, 3: 14, 4: 14},
        nom_tableau="Catalogue",
    )

    wb.create_sheet("Saisies")
    core.ecrire_feuille(
        wb["Saisies"],
        ["ID_Ligne", "Date", "Code_Produit", "Quantite_Vendue",
         "CA_Realise", "CA_Objectif",
         "Prix_Vente_Catalogue (à compléter)",
         "Marge_Unitaire (à compléter)",
         "Taux_Atteinte_Objectif (à compléter)"],
        saisies,
        formats={2: "DD/MM/YYYY", 5: "# ##0.00 €", 6: "# ##0.00 €",
                 7: "# ##0.00 €", 8: "# ##0.00 €", 9: "0.0 %"},
        largeurs={1: 10, 2: 12, 3: 14, 4: 16, 5: 14, 6: 14, 7: 22, 8: 20,
                  9: 24},
        nom_tableau="Saisies",
    )
    for lettre in ("G", "H", "I"):
        wb["Saisies"]["%s1" % lettre].fill = JAUNE
        wb["Saisies"]["%s1" % lettre].font = Font(bold=True, color="7F6000")

    # ------------------------------------------------------- recherche
    ws = wb.create_sheet("Recherche")
    ws["A1"] = "Zone de recherche par code produit"
    ws["A1"].font = Font(bold=True, size=13, color="1F3864")

    ws["A3"] = "Code produit recherché :"
    ws["B3"] = "P003"
    ws["B3"].fill = JAUNE
    ws["B3"].font = Font(bold=True)

    ws["A5"] = "Désignation (à compléter) :"
    ws["A6"] = "Prix de vente (à compléter) :"
    ws["A7"] = "Taux de marge (à compléter) :"

    ws["A9"] = "Codes de test à utiliser en B3 :"
    ws["A10"] = "P003"
    ws["B10"] = "code présent au catalogue"
    ws["A11"] = "P105"
    ws["B11"] = "présent, mais prix de vente à 0"
    ws["A12"] = "P777"
    ws["B12"] = "code absent du catalogue"
    for ligne in range(9, 13):
        ws["A%d" % ligne].alignment = Alignment(horizontal="left")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 34
    ws.sheet_view.showGridLines = False

    core.enregistrer(wb, "Suivi_Marges.xlsx")
    print("     Catalogue : %d references (catalogue complet)"
          % len(catalogue))
    print("     Saisies   : %d lignes (9 #N/A, 13 #DIV/0! potentiels)"
          % len(saisies))


if __name__ == "__main__":
    main()
